import docx
import openpyxl
from openpyxl.styles import Alignment

def convert_word_table_to_excel(word_path, excel_path):
    # 打开Word文档
    doc = docx.Document(word_path)
    
    # 创建一个新的Excel工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # 遍历Word文档中的所有表格
    for table in doc.tables:
        # 遍历表格的每一行
        for row in table.rows:
            # 提取每一行的单元格内容
            row_data = [cell.text.strip() for cell in row.cells]
            # 将数据写入Excel
            ws.append(row_data)
    
    # 设置单元格对齐方式
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 保存Excel文件
    wb.save(excel_path)
    print(f"文件已保存到：{excel_path}")

# 使用示例
word_path = "澳斯卡表格.docx"  # 替换为你的Word文件路径
excel_path = "/Users/dabaobaodemac/Desktop/aaa.xlsx"    # 替换为你希望保存的Excel文件路径
convert_word_table_to_excel(word_path, excel_path)